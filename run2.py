from openpyxl import load_workbook
from docx import Document
from docx.oxml.ns import qn
from datetime import datetime
from pathlib import Path
import tempfile
import shutil
import unicodedata
import sys
import re

if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "Dinamicna_polja_SE_BHEE.xlsx"
WORD_FILE = BASE_DIR / "Analiza_osnovna.docx"
OUTPUT_FILE = BASE_DIR / "Analiza_izpolnjena.docx"


def format_number(value):
    if value is None:
        return ""

    if not isinstance(value, (int, float)):
        return str(value)

    formatted = f"{value:,.2f}"
    formatted = formatted.replace(",", "TEMP")
    formatted = formatted.replace(".", ",")
    formatted = formatted.replace("TEMP", ".")

    return formatted


def format_value(value, unit_type=""):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")

    unit_type = str(unit_type or "").strip()

    if isinstance(value, (int, float)):
        value = format_number(value)
    else:
        value = str(value)

    units_to_append = {
        "EUR",
        "kWp",
        "kW",
        "kWh",
        "EUR/MWh",
        "EUR/leto",
        "kWh/leto",
        "kWh/kWp/leto",
        "leta"
    }

    if unit_type in units_to_append:
        return f"{value} {unit_type}"

    return value


def make_temp_copy(source_file, temp_dir):
    destination = Path(temp_dir) / source_file.name
    shutil.copy2(source_file, destination)
    return destination


def normalize_compare_value(value):
    if value is None:
        return ""

    value = str(value).strip().lower()

    value = unicodedata.normalize(
        "NFKD",
        value
    )

    value = "".join(
        char
        for char in value
        if not unicodedata.combining(char)
    )

    value = re.sub(
        r"\s+",
        " ",
        value
    )

    return value.strip()


def read_excel_fields(filename):
    workbook = load_workbook(
        filename,
        data_only=True
    )

    worksheet = workbook.active

    headers = {}

    for cell in worksheet[1]:
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column

    required = [
        "ID_polja",
        "Excel_ime_polja",
        "Vrednost"
    ]

    for column in required:
        if column not in headers:
            workbook.close()

            raise ValueError(
                f"V Excelu manjka stolpec: {column}"
            )

    fields = {}
    dp_values = {}
    dp_units = {}

    for row in range(
        2,
        worksheet.max_row + 1
    ):
        field_id = worksheet.cell(
            row,
            headers["ID_polja"]
        ).value

        excel_name = worksheet.cell(
            row,
            headers["Excel_ime_polja"]
        ).value

        value = worksheet.cell(
            row,
            headers["Vrednost"]
        ).value

        unit_type = ""

        if "Enota_tip" in headers:
            unit_type = worksheet.cell(
                row,
                headers["Enota_tip"]
            ).value

        if not field_id or not excel_name:
            continue

        field_id = str(field_id).strip()
        excel_name = str(excel_name).strip()
        unit_type = str(unit_type or "").strip()

        placeholder = f"[{field_id}_{excel_name}]"

        fields[placeholder] = format_value(
            value,
            unit_type
        )

        dp_values[field_id.upper()] = value
        dp_units[field_id.upper()] = unit_type

    workbook.close()

    return (
        fields,
        dp_values,
        dp_units
    )


def numeric_value(value):
    if isinstance(value, (int, float)):
        return float(value)

    value_string = (
        str(value)
        .strip()
        .replace(" ", "")
    )

    if "." in value_string and "," in value_string:
        value_string = (
            value_string
            .replace(".", "")
            .replace(",", ".")
        )

    elif "," in value_string:
        value_string = value_string.replace(",", ".")

    return float(value_string)


def calculate_expression(
    expression,
    dp_values
):
    dp_ids = re.findall(
        r"DP\d+",
        expression,
        flags=re.IGNORECASE
    )

    if len(dp_ids) < 2:
        return None

    python_expression = expression

    for dp_id in dp_ids:
        lookup_id = dp_id.upper()

        if lookup_id not in dp_values:
            return None

        value = dp_values[lookup_id]

        if value is None:
            return None

        try:
            value = numeric_value(value)

        except ValueError:
            return None

        python_expression = re.sub(
            re.escape(dp_id),
            str(value),
            python_expression,
            count=1,
            flags=re.IGNORECASE
        )

    if not re.fullmatch(
        r"[0-9+\-*/().\s]+",
        python_expression
    ):
        return None

    try:
        result = eval(
            python_expression,
            {"__builtins__": {}},
            {}
        )

    except (
        ZeroDivisionError,
        SyntaxError,
        TypeError
    ):
        return None

    return result


def get_expression_unit(
    expression,
    dp_units
):
    dp_ids = re.findall(
        r"DP\d+",
        expression,
        flags=re.IGNORECASE
    )

    if not dp_ids:
        return ""

    units = [
        dp_units.get(
            dp_id.upper(),
            ""
        )
        for dp_id in dp_ids
    ]

    if "+" in expression or "-" in expression:
        if (
            units
            and all(
                unit == units[0]
                for unit in units
            )
        ):
            return units[0]

    return ""


def replace_expression_placeholders(
    text,
    dp_values,
    dp_units
):
    pattern = re.compile(
        r"\[((?:DP\d+)"
        r"(?:[+\-*/](?:DP\d+))+)"
        r"(?:_[^\]]+)?\]",
        flags=re.IGNORECASE
    )

    def replacement(match):
        expression = match.group(1)

        result = calculate_expression(
            expression,
            dp_values
        )

        if result is None:
            return match.group(0)

        unit = get_expression_unit(
            expression,
            dp_units
        )

        return format_value(
            result,
            unit
        )

    return pattern.sub(
        replacement,
        text
    )


def get_element_text(element):
    text_parts = []

    for text_element in element.iter(
        qn("w:t")
    ):
        if text_element.text:
            text_parts.append(
                text_element.text
            )

    return "".join(
        text_parts
    ).strip()


def parse_conditional_marker(text):
    text = text.strip()

    match = re.fullmatch(
        r"\[(START|END)_(DP\d+)_(.+?)\]",
        text,
        flags=re.IGNORECASE
    )

    if not match:
        return None

    return {
        "type": match.group(1).upper(),
        "dp_id": match.group(2).upper(),
        "expected_value": match.group(3).strip()
    }


def find_conditional_blocks(document):
    body = document._element.body
    elements = list(body.iterchildren())

    blocks = []
    current_block = None

    for index, element in enumerate(elements):
        text = get_element_text(element)
        marker = parse_conditional_marker(text)

        if marker is None:
            continue

        if marker["type"] == "START":
            if current_block is not None:
                raise ValueError(
                    "Pogojni bloki se ne smejo "
                    "prekrivati ali gnezditi. "
                    f"Nov START marker: {text}"
                )

            current_block = {
                "dp_id": marker["dp_id"],
                "expected_value": marker[
                    "expected_value"
                ],
                "start_index": index,
                "start_text": text
            }

        elif marker["type"] == "END":
            if current_block is None:
                raise ValueError(
                    f"Najden END brez START: {text}"
                )

            if (
                current_block["dp_id"]
                != marker["dp_id"]
            ):
                raise ValueError(
                    "START in END nimata "
                    "istega DP polja."
                )

            if (
                normalize_compare_value(
                    current_block[
                        "expected_value"
                    ]
                )
                !=
                normalize_compare_value(
                    marker[
                        "expected_value"
                    ]
                )
            ):
                raise ValueError(
                    "START in END nimata "
                    "iste pričakovane vrednosti."
                )

            current_block["end_index"] = index
            current_block["end_text"] = text

            blocks.append(current_block)
            current_block = None

    if current_block is not None:
        raise ValueError(
            "Za START marker manjka END: "
            f"{current_block['start_text']}"
        )

    return blocks


def remove_xml_element(element):
    parent = element.getparent()

    if parent is not None:
        parent.remove(element)


def clear_xml_element_text(element):
    for text_element in element.iter(
        qn("w:t")
    ):
        text_element.text = ""


def process_conditional_sections(
    document,
    dp_values
):
    body = document._element.body

    blocks = find_conditional_blocks(
        document
    )

    if not blocks:
        print(
            "V Wordu ni pogojnih "
            "START/END blokov."
        )
        return

    print()
    print(
        f"Najdenih pogojnih blokov: "
        f"{len(blocks)}"
    )

    elements = list(
        body.iterchildren()
    )

    for block in reversed(blocks):
        dp_id = block["dp_id"]
        expected_value = block[
            "expected_value"
        ]

        start_index = block[
            "start_index"
        ]

        end_index = block[
            "end_index"
        ]

        if dp_id not in dp_values:
            clear_xml_element_text(
                elements[start_index]
            )

            clear_xml_element_text(
                elements[end_index]
            )

            print(
                f"{dp_id} ne obstaja v Excelu -> "
                f"blok '{expected_value}' OSTANE, "
                f"markerja odstranjena"
            )

            continue

        excel_value = dp_values[
            dp_id
        ]

        expected_normalized = (
            normalize_compare_value(
                expected_value
            )
        )

        actual_normalized = (
            normalize_compare_value(
                excel_value
            )
        )

        if (
            actual_normalized
            == expected_normalized
        ):
            clear_xml_element_text(
                elements[start_index]
            )

            clear_xml_element_text(
                elements[end_index]
            )

            print(
                f"{dp_id} = {excel_value} -> "
                f"'{expected_value}' OSTANE"
            )

        else:
            for index in range(
                end_index,
                start_index - 1,
                -1
            ):
                remove_xml_element(
                    elements[index]
                )

            print(
                f"{dp_id} = {excel_value} -> "
                f"'{expected_value}' IZBRISAN"
            )


def replace_in_paragraph(
    paragraph,
    fields,
    dp_values,
    dp_units
):
    full_text = "".join(
        run.text
        for run in paragraph.runs
    )

    if not full_text:
        return

    new_text = (
        replace_expression_placeholders(
            full_text,
            dp_values,
            dp_units
        )
    )

    for placeholder, value in fields.items():
        new_text = new_text.replace(
            placeholder,
            value
        )

    if new_text == full_text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = new_text

        for run in paragraph.runs[1:]:
            run.text = ""


def replace_in_table(
    table,
    fields,
    dp_values,
    dp_units
):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(
                    paragraph,
                    fields,
                    dp_values,
                    dp_units
                )

            for nested_table in cell.tables:
                replace_in_table(
                    nested_table,
                    fields,
                    dp_values,
                    dp_units
                )


def replace_in_document(
    document,
    fields,
    dp_values,
    dp_units
):
    for paragraph in document.paragraphs:
        replace_in_paragraph(
            paragraph,
            fields,
            dp_values,
            dp_units
        )

    for table in document.tables:
        replace_in_table(
            table,
            fields,
            dp_values,
            dp_units
        )

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(
                paragraph,
                fields,
                dp_values,
                dp_units
            )

        for table in section.header.tables:
            replace_in_table(
                table,
                fields,
                dp_values,
                dp_units
            )

        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(
                paragraph,
                fields,
                dp_values,
                dp_units
            )

        for table in section.footer.tables:
            replace_in_table(
                table,
                fields,
                dp_values,
                dp_units
            )


def find_remaining_placeholders(document):
    pattern = re.compile(
        r"\[DP[^\]]+\]"
    )

    found = set()

    def check_paragraph(paragraph):
        text = "".join(
            run.text
            for run in paragraph.runs
        )

        for match in pattern.findall(
            text
        ):
            found.add(match)

    def check_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    check_paragraph(
                        paragraph
                    )

                for nested_table in cell.tables:
                    check_table(
                        nested_table
                    )

    for paragraph in document.paragraphs:
        check_paragraph(
            paragraph
        )

    for table in document.tables:
        check_table(
            table
        )

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            check_paragraph(
                paragraph
            )

        for table in section.header.tables:
            check_table(
                table
            )

        for paragraph in section.footer.paragraphs:
            check_paragraph(
                paragraph
            )

        for table in section.footer.tables:
            check_table(
                table
            )

    return found


def find_remaining_conditional_markers(
    document
):
    found = set()

    body = document._element.body

    for element in body.iterchildren():
        text = get_element_text(
            element
        )

        if parse_conditional_marker(
            text
        ):
            found.add(
                text
            )

    return found


def save_document(document):
    try:
        document.save(
            OUTPUT_FILE
        )

        return OUTPUT_FILE

    except PermissionError:
        timestamp = (
            datetime.now()
            .strftime(
                "%Y%m%d_%H%M%S"
            )
        )

        alternative = (
            BASE_DIR
            / (
                "Analiza_izpolnjena_"
                f"{timestamp}.docx"
            )
        )

        document.save(
            alternative
        )

        print()
        print(
            "POZOR: Analiza_izpolnjena.docx "
            "je trenutno odprta."
        )

        print(
            "Dokument je shranjen kot:"
        )

        print(
            alternative.name
        )

        return alternative


def main():
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(
            f"Manjka datoteka: "
            f"{EXCEL_FILE.name}"
        )

    if not WORD_FILE.exists():
        raise FileNotFoundError(
            f"Manjka datoteka: "
            f"{WORD_FILE.name}"
        )

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_excel = make_temp_copy(
            EXCEL_FILE,
            temp_dir
        )

        temp_word = make_temp_copy(
            WORD_FILE,
            temp_dir
        )

        (
            fields,
            dp_values,
            dp_units
        ) = read_excel_fields(
            temp_excel
        )

        print(
            f"Prebranih polj iz Excela: "
            f"{len(fields)}"
        )

        document = Document(
            temp_word
        )

        process_conditional_sections(
            document,
            dp_values
        )

        replace_in_document(
            document,
            fields,
            dp_values,
            dp_units
        )

        remaining = (
            find_remaining_placeholders(
                document
            )
        )

        remaining_markers = (
            find_remaining_conditional_markers(
                document
            )
        )

        if remaining_markers:
            raise ValueError(
                "V dokumentu so ostali "
                "START/END markerji: "
                + ", ".join(
                    sorted(
                        remaining_markers
                    )
                )
            )

        saved_file = save_document(
            document
        )

    print()
    print(
        f"Dokument shranjen: "
        f"{saved_file.name}"
    )

    if remaining:
        print()

        print(
            "POZOR: naslednja polja "
            "niso bila zamenjana:"
        )

        for placeholder in sorted(
            remaining
        ):
            print(
                f"  {placeholder}"
            )

    else:
        print(
            "Vsa dinamična polja "
            "so bila zamenjana."
        )


if __name__ == "__main__":
    try:
        main()

    except PermissionError:
        print()

        print(
            "NAPAKA: Datoteke trenutno "
            "ni mogoče prebrati."
        )

        print(
            "Preveri OneDrive sinhronizacijo."
        )

        print()

        input(
            "Pritisni Enter za izhod..."
        )

    except Exception as e:
        print()

        print(
            f"NAPAKA: {e}"
        )

        print()

        input(
            "Pritisni Enter za izhod..."
        )