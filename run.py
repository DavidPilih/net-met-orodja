from openpyxl import load_workbook
from docx import Document
from datetime import datetime
from pathlib import Path
import sys
import re

if getattr(sys, "frozen", False):
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "Dinamicna_polja_SE_BHEE.xlsx"
WORD_FILE = BASE_DIR / "Analiza_osnovna.docx"
OUTPUT_FILE = BASE_DIR / "Analiza_izpolnjena.docx"


def format_number(value):
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))

        return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", ",")

    return str(value)


def format_value(value, unit_type=""):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")

    unit_type = str(unit_type or "").strip()
    value = format_number(value)

    units_to_append = {
        "EUR",
        "kWp",
        "kW",
        "kWh",
        "EUR/MWh",
        "EUR/leto",
        "kWh/leto",
        "kWh/kWp/leto",
        "leta"
    }

    if unit_type in units_to_append:
        return f"{value} {unit_type}"

    return value


def read_excel_fields(filename):
    workbook = load_workbook(filename, data_only=True)
    worksheet = workbook.active

    headers = {}

    for cell in worksheet[1]:
        if cell.value is not None:
            headers[str(cell.value).strip()] = cell.column

    required = ["ID_polja", "Excel_ime_polja", "Vrednost"]

    for column in required:
        if column not in headers:
            raise ValueError(f"V Excelu manjka stolpec: {column}")

    fields = {}
    dp_values = {}
    dp_units = {}

    for row in range(2, worksheet.max_row + 1):
        field_id = worksheet.cell(row, headers["ID_polja"]).value
        excel_name = worksheet.cell(row, headers["Excel_ime_polja"]).value
        value = worksheet.cell(row, headers["Vrednost"]).value

        unit_type = ""

        if "Enota_tip" in headers:
            unit_type = worksheet.cell(row, headers["Enota_tip"]).value

        if not field_id or not excel_name:
            continue

        field_id = str(field_id).strip()
        excel_name = str(excel_name).strip()
        unit_type = str(unit_type or "").strip()

        placeholder = f"[{field_id}_{excel_name}]"

        fields[placeholder] = format_value(value, unit_type)
        dp_values[field_id] = value
        dp_units[field_id] = unit_type

    return fields, dp_values, dp_units


def calculate_expression(expression, dp_values):
    dp_ids = re.findall(r"DP\d+", expression)

    if len(dp_ids) < 2:
        return None

    python_expression = expression

    for dp_id in dp_ids:
        if dp_id not in dp_values:
            return None

        value = dp_values[dp_id]

        if value is None:
            return None

        if not isinstance(value, (int, float)):
            try:
                value = float(str(value).replace(",", "."))
            except ValueError:
                return None

        python_expression = python_expression.replace(dp_id, str(value))

    if not re.fullmatch(r"[0-9+\-*/().\s]+", python_expression):
        return None

    try:
        result = eval(python_expression, {"__builtins__": {}}, {})
    except (ZeroDivisionError, SyntaxError, TypeError):
        return None

    return result


def get_expression_unit(expression, dp_units):
    dp_ids = re.findall(r"DP\d+", expression)

    if not dp_ids:
        return ""

    units = [dp_units.get(dp_id, "") for dp_id in dp_ids]

    if "+" in expression or "-" in expression:
        if units and all(unit == units[0] for unit in units):
            return units[0]

    return ""


def replace_expression_placeholders(text, dp_values, dp_units):
    pattern = re.compile(
        r"\[((?:DP\d+)(?:[+\-*/](?:DP\d+))+)(?:_[^\]]+)?\]"
    )

    def replacement(match):
        expression = match.group(1)

        result = calculate_expression(expression, dp_values)

        if result is None:
            return match.group(0)

        unit = get_expression_unit(expression, dp_units)

        return format_value(result, unit)

    return pattern.sub(replacement, text)


def replace_in_paragraph(paragraph, fields, dp_values, dp_units):
    full_text = "".join(run.text for run in paragraph.runs)

    if not full_text:
        return

    new_text = full_text

    new_text = replace_expression_placeholders(
        new_text,
        dp_values,
        dp_units
    )

    for placeholder, value in fields.items():
        new_text = new_text.replace(placeholder, value)

    if new_text == full_text:
        return

    if paragraph.runs:
        paragraph.runs[0].text = new_text

        for run in paragraph.runs[1:]:
            run.text = ""


def replace_in_table(table, fields, dp_values, dp_units):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(
                    paragraph,
                    fields,
                    dp_values,
                    dp_units
                )

            for nested_table in cell.tables:
                replace_in_table(
                    nested_table,
                    fields,
                    dp_values,
                    dp_units
                )


def replace_in_document(document, fields, dp_values, dp_units):
    for paragraph in document.paragraphs:
        replace_in_paragraph(
            paragraph,
            fields,
            dp_values,
            dp_units
        )

    for table in document.tables:
        replace_in_table(
            table,
            fields,
            dp_values,
            dp_units
        )

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            replace_in_paragraph(
                paragraph,
                fields,
                dp_values,
                dp_units
            )

        for table in section.header.tables:
            replace_in_table(
                table,
                fields,
                dp_values,
                dp_units
            )

        for paragraph in section.footer.paragraphs:
            replace_in_paragraph(
                paragraph,
                fields,
                dp_values,
                dp_units
            )

        for table in section.footer.tables:
            replace_in_table(
                table,
                fields,
                dp_values,
                dp_units
            )


def find_remaining_placeholders(document):
    pattern = re.compile(r"\[DP[^\]]+\]")
    found = set()

    def check_paragraph(paragraph):
        text = "".join(run.text for run in paragraph.runs)

        for match in pattern.findall(text):
            found.add(match)

    def check_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    check_paragraph(paragraph)

                for nested_table in cell.tables:
                    check_table(nested_table)

    for paragraph in document.paragraphs:
        check_paragraph(paragraph)

    for table in document.tables:
        check_table(table)

    for section in document.sections:
        for paragraph in section.header.paragraphs:
            check_paragraph(paragraph)

        for table in section.header.tables:
            check_table(table)

        for paragraph in section.footer.paragraphs:
            check_paragraph(paragraph)

        for table in section.footer.tables:
            check_table(table)

    return found


def main():
    fields, dp_values, dp_units = read_excel_fields(EXCEL_FILE)

    print(f"Prebranih polj iz Excela: {len(fields)}")

    for placeholder, value in fields.items():
        print(f"{placeholder} -> {value}")

    document = Document(WORD_FILE)

    replace_in_document(
        document,
        fields,
        dp_values,
        dp_units
    )

    remaining = find_remaining_placeholders(document)

    document.save(OUTPUT_FILE)

    print()
    print(f"Dokument shranjen: {OUTPUT_FILE}")

    if remaining:
        print()
        print("POZOR: naslednja Word polja niso bila zamenjana:")

        for placeholder in sorted(remaining):
            print(f"  {placeholder}")

    else:
        print("Vsa dinamična polja so bila zamenjana.")


if __name__ == "__main__":
    try:
        main()

    except PermissionError:
        print()
        print("NAPAKA: Datoteka je verjetno odprta v Excelu ali Wordu.")
        print("Prosim zapri Excel in Word ter poskusi ponovno.")
        print()
        input("Pritisni Enter za izhod...")

    except Exception as e:
        print()
        print(f"NAPAKA: {e}")
        print()
        input("Pritisni Enter za izhod...")